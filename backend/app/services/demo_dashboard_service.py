from __future__ import annotations

import json
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from backend.app.config import ROOT_DIR
from backend.app.services.demo_accounts import DEMO_USERS


DEMO_WORKBOOK_PATH = ROOT_DIR / "data" / "demo_process_kpi_data.xlsx"
DEMO_JSON_PATH = DEMO_WORKBOOK_PATH.with_suffix(".json")
DEMO_ACCOUNT_ASSIGNMENTS = {
    user["email"]: {
        "role": user["role"],
        "persona_id": user["persona_id"],
        "designation": f"Role {int(user['persona_id'].removeprefix('role-')):02d} - {user['persona_name']}",
    }
    for user in DEMO_USERS
}
DEMO_ACCOUNT_ASSIGNMENTS.update({
    "admin@local.dev": {
        "role": "operations",
        "persona_id": "role-09",
        "designation": "Role 09 - Operations Manager",
    },
    "sales@local.dev": {
        "role": "sales",
        "persona_id": "role-14",
        "designation": "Role 14 - Sales Rep",
    },
})


def _read_table(workbook, sheet_name: str) -> list[dict[str, Any]]:
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    return [
        dict(zip(headers, row))
        for row in rows
        if any(value is not None for value in row)
    ]


@lru_cache(maxsize=2)
def _load_demo_json(path: str, modified_ns: int, file_size: int) -> dict[str, list[dict[str, Any]]]:
    del modified_ns, file_size
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _read_demo_table(
    sheet_name: str,
    workbook_path: Path = DEMO_WORKBOOK_PATH,
) -> list[dict[str, Any]]:
    json_path = DEMO_JSON_PATH if workbook_path == DEMO_WORKBOOK_PATH else workbook_path.with_suffix(".json")
    if workbook_path == DEMO_WORKBOOK_PATH and json_path.exists():
        stat = json_path.stat()
        tables = _load_demo_json(str(json_path), stat.st_mtime_ns, stat.st_size)
        return tables[sheet_name]
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        return _read_table(workbook, sheet_name)
    finally:
        workbook.close()


def _number(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _measurement_value(measurement: dict[str, Any], definition: dict[str, Any]) -> float:
    aggregation = definition["aggregation_method"]
    numerator = _number(measurement.get("numerator"))
    denominator = _number(measurement.get("denominator"))
    if aggregation == "ratio_percent":
        return round(numerator / denominator * 100, 2) if denominator else 0.0
    if aggregation == "forecast_accuracy_percent":
        return round((1 - numerator / denominator) * 100, 2) if denominator else 0.0
    if aggregation == "average":
        return round(numerator / denominator, 2) if denominator else 0.0
    return round(numerator, 2)


def _period_bounds(rows: list[dict[str, Any]], period: str) -> tuple[datetime | None, datetime | None]:
    dates = [
        datetime.fromisoformat(str(row["process_start_timestamp"]).replace("Z", "+00:00"))
        for row in rows
        if row.get("process_start_timestamp")
    ]
    if not dates:
        return None, None
    latest = max(dates)
    if period == "daily":
        start = latest.replace(hour=0, minute=0, second=0, microsecond=0)
        return start, start.replace(hour=23, minute=59, second=59, microsecond=999999)
    if period == "weekly":
        start = (latest - timedelta(days=latest.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
        end = (start + timedelta(days=7)).replace(microsecond=0) - timedelta(microseconds=1)
        return start, end
    if period == "monthly":
        start = latest.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        next_month = start.replace(year=start.year + (start.month == 12), month=(start.month % 12) + 1)
        return start, next_month - timedelta(microseconds=1)
    if period == "quarterly":
        month = ((latest.month - 1) // 3) * 3 + 1
        start = latest.replace(month=month, day=1, hour=0, minute=0, second=0, microsecond=0)
        if month == 10:
            next_quarter = start.replace(year=start.year + 1, month=1)
        else:
            next_quarter = start.replace(month=month + 3)
        return start, next_quarter - timedelta(microseconds=1)
    start = latest.replace(month=1, day=1, hour=0, minute=0, second=0, microsecond=0)
    return start, start.replace(year=start.year + 1) - timedelta(microseconds=1)


def _definition_components(definition: dict[str, Any], aggregation: str) -> tuple[str | None, str | None]:
    numerator = definition.get("numerator_label")
    denominator = definition.get("denominator_label")
    if aggregation == "derived_from_executions" and "mean duration" in str(definition.get("formula", "")).lower():
        return "Total completed duration (minutes)", "Completed executions"
    return numerator, denominator


def _status_for_target(value: float | None, target: float | None, direction: str | None) -> str:
    if value is None:
        return "No data"
    if target is None:
        return "No target configured"
    if direction and direction.lower() == "increase":
        return "On target" if value >= target else "Below target"
    if direction and direction.lower() == "decrease":
        return "On target" if value <= target else "Above target"
    return "Target rule not configured"


def _build_kpi_detail(
    definition: dict[str, Any],
    measurements: list[dict[str, Any]],
    observations: list[dict[str, Any]],
    executions: list[dict[str, Any]],
    period: str,
    last_updated: str,
    persona_id: str,
    persona_name: str,
    viewer_name: str,
    record_limit: int,
    record_offset: int,
) -> dict[str, Any]:
    aggregation = str(definition.get("aggregation_method", ""))
    is_execution_kpi = aggregation == "derived_from_executions"
    available_period = str(definition.get("measurement_period", "weekly")).lower()
    points: list[dict[str, Any]] = []
    numerator: float | None = None
    denominator: float | None = None
    sample_size = 0
    records: list[dict[str, Any]] = []
    total_observation_count = len(observations)
    trend_series: list[dict[str, Any]] = []
    team_average: float | None = None

    if observations:
        dated_observations = [
            row for row in observations if row.get("observed_at")
        ]
        start, end = _period_bounds([
            {"process_start_timestamp": row["observed_at"]}
            for row in dated_observations
        ], period)
        period_observations = [
            row for row in dated_observations
            if start is not None
            and end is not None
            and start <= datetime.fromisoformat(str(row["observed_at"]).replace("Z", "+00:00")) <= end
        ]
        available = bool(period_observations)
        numerator = sum(_number(row.get("numerator")) for row in period_observations) if available else None
        denominator = sum(_number(row.get("denominator")) for row in period_observations) if available else None
        value = _measurement_value(
            {"numerator": numerator, "denominator": denominator},
            definition,
        ) if available and (
            denominator
            or definition.get("aggregation_method") == "count"
        ) else None
        sample_size = len(period_observations)
        period_label = period.title()
        if period_observations:
            period_start = min(str(row["observed_at"]) for row in period_observations)
            period_end = max(str(row["observed_at"]) for row in period_observations)
            sorted_observations = sorted(
                period_observations,
                key=lambda row: str(row["observed_at"]),
                reverse=True,
            )
            records = [
                {
                    "observation_id": row["observation_id"],
                    "account_name": row["synthetic_account_name"],
                    "account_id": row["synthetic_account_id"],
                    "observed_at": str(row["observed_at"]),
                    "component_a_label": row["component_a_label"],
                    "component_a_value": _number(row["component_a_value"]),
                    "component_b_label": row["component_b_label"],
                    "component_b_value": _number(row["component_b_value"]),
                    "numerator": _number(row["numerator"]),
                    "denominator": _number(row["denominator"]),
                    "result_label": row["result_label"],
                    "result_value": _number(row["result_value"]),
                }
                for row in sorted_observations[record_offset : record_offset + record_limit]
            ]
        else:
            period_start = None
            period_end = None
    elif is_execution_kpi:
        start, end = _period_bounds(executions, period)
        period_executions = [
            row for row in executions
            if start is not None
            and end is not None
            and row.get("process_execution_status") != "Pending"
            and start <= datetime.fromisoformat(str(row["process_start_timestamp"]).replace("Z", "+00:00")) <= end
        ]
        completed = [row for row in period_executions if row["process_execution_status"] == "Completed"]
        if definition["kpi_id"] == "process-completion-rate":
            numerator, denominator = float(len(completed)), float(len(period_executions))
            value = round(numerator / denominator * 100, 2) if denominator else None
        elif definition["kpi_id"] == "average-process-duration":
            durations = [
                _number(row["actual_duration_minutes"])
                for row in completed if row.get("actual_duration_minutes") is not None
            ]
            numerator, denominator = sum(durations), float(len(durations))
            value = round(numerator / denominator, 2) if denominator else None
        elif definition["kpi_id"] == "on-time-completion-rate":
            on_time = [
                row for row in completed
                if _number(row.get("actual_duration_minutes")) <= _number(row.get("sla_target_duration_minutes"))
            ]
            numerator, denominator = float(len(on_time)), float(len(completed))
            value = round(numerator / denominator * 100, 2) if denominator else None
        else:
            value = None
        sample_size = len(period_executions)
        available = value is not None
        period_label = period.title()
    else:
        available = period == available_period and bool(measurements)
        points = sorted(measurements, key=lambda row: str(row.get("period_end", "")))
        if available:
            latest = points[-1]
            numerator = _number(latest.get("numerator")) if latest.get("numerator") is not None else None
            denominator = _number(latest.get("denominator")) if latest.get("denominator") is not None else None
            value = _measurement_value(latest, definition)
            sample_size = int(latest.get("sample_size") or 0)
        else:
            value = None
        period_label = period.title()

    target_value = definition.get("target_value")
    target = _number(target_value) if target_value is not None else None
    if observations:
        team_average = round(
            sum(_number(row.get("result_value")) for row in observations) / len(observations),
            2,
        )
        weekly_points = sorted(measurements, key=lambda row: str(row.get("period_end", "")))
        for point in weekly_points[-8:]:
            trend_series.append({
                "time": str(point.get("period_start", "")),
                "value": _measurement_value(point, definition),
            })
    formula = definition.get("formula") or definition.get("description")
    numerator_label, denominator_label = _definition_components(definition, aggregation)
    components = []
    if numerator_label:
        components.append({"label": numerator_label, "value": numerator, "component": "numerator"})
    if denominator_label:
        components.append({"label": denominator_label, "value": denominator, "component": "denominator"})
    if value is not None:
        components.append({"label": definition["kpi_name"], "value": value, "unit": definition["measurement_unit"], "component": "result"})

    if observations:
        values = [point["value"] for point in trend_series]
        if len(values) > 1:
            trend = "up" if values[-1] > values[-2] else "down" if values[-1] < values[-2] else "stable"
        else:
            trend = "Not available"
    elif is_execution_kpi:
        trend = "Not available"
        period_start = None
        period_end = None
    elif available:
        values = [_measurement_value(point, definition) for point in points]
        if len(values) > 1:
            trend = "up" if values[-1] > values[-2] else "down" if values[-1] < values[-2] else "stable"
        else:
            trend = "Not available"
        period_start = points[-1].get("period_start")
        period_end = points[-1].get("period_end")
    else:
        trend = "Not available"
        period_start = None
        period_end = None

    return {
        "kpi_id": definition["kpi_id"],
        "kpi_name": definition["kpi_name"],
        "name": definition["kpi_name"],
        "description": definition.get("description"),
        "definition": definition.get("definition") or definition.get("description"),
        "formula": formula,
        "unit": definition.get("measurement_unit"),
        "target": target,
        "target_source": definition.get("target_source"),
        "current_value": value,
        "team_average": team_average,
        "team_average_variance": round(value - team_average, 2) if value is not None and team_average is not None else None,
        "team_average_distance": round(abs(value - team_average), 2) if value is not None and team_average is not None else None,
        "team_average_label": "Synthetic team average across all 150 demo-account observations" if observations else None,
        "trend_series": trend_series,
        "value": value,
        "variance_to_target": round(value - target, 2) if value is not None and target is not None else None,
        "status": (
            f"{_status_for_target(value, target, definition.get('direction'))} (illustrative)"
            if definition.get("target_source") and _status_for_target(value, target, definition.get("direction")) in {"On target", "Below target", "Above target"}
            else _status_for_target(value, target, definition.get("direction"))
        ),
        "trend": trend,
        "numerator": numerator,
        "denominator": denominator,
        "numerator_label": numerator_label,
        "denominator_label": denominator_label,
        "components": components,
        "sample_size": sample_size,
        "observation_count": sample_size,
        "total_observation_count": total_observation_count,
        "records": records,
        "record_offset": record_offset,
        "record_limit": record_limit,
        "period": period_label,
        "period_start": str(period_start) if period_start is not None else None,
        "period_end": str(period_end) if period_end is not None else None,
        "data_available": available,
        "availability_message": None if available else "Demo data not available for this period",
        "persona": f"{persona_id.replace('role-', 'Role ')} - {persona_name}",
        "persona_id": persona_id,
        "viewer_name": viewer_name,
        "driver": definition.get("driver"),
        "enterprise_linkage": definition.get("enterprise_linkage"),
        "function": definition.get("function"),
        "process": definition.get("process"),
        "activity": definition.get("activity"),
        "task": definition.get("task"),
        "data_source": "DEMO / SYNTHETIC DATA",
        "data_source_detail": (
            "Authenticated FastAPI API backed by KPI_Definitions and KPI_Observations "
            "in the generated demo_process_kpi_data.xlsx workbook"
            if observations
            else "Authenticated FastAPI API backed by KPI_Definitions, KPI_Measurements, "
            "and Process_Executions in the generated demo_process_kpi_data.xlsx workbook"
        ),
        "source": definition.get("source") or definition.get("data_source_type"),
        "api_endpoint": f"/api/kpis/{definition['kpi_id']}/details?period={period}",
        "last_updated": last_updated,
    }


def list_demo_kpis(
    email: str,
    persona_id: str,
    period: str,
    workbook_path: Path = DEMO_WORKBOOK_PATH,
    record_limit: int = 10,
    record_offset: int = 0,
) -> list[dict[str, Any]]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Demo KPI workbook not found: {workbook_path}")

    definitions = _read_demo_table("KPI_Definitions", workbook_path)
    measurements = _read_demo_table("KPI_Measurements", workbook_path)
    executions = _read_demo_table("Process_Executions", workbook_path)
    persona_resources = _read_demo_table("Persona_Resource_Mapping", workbook_path)
    observations = _read_demo_table("KPI_Observations", workbook_path)

    normalized_email = email.strip().lower()
    user_measurements: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in measurements:
        if str(row.get("resource_email", "")).lower() == normalized_email:
            user_measurements[str(row["kpi_id"])].append(row)
    user_executions = [
        row for row in executions
        if str(row.get("assigned_resource_email", "")).lower() == normalized_email
    ]
    user_observations: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in observations:
        if row.get("persona_id") == persona_id:
            user_observations[str(row["kpi_id"])].append(row)
    resource = next(
        (row for row in persona_resources if str(row.get("email", "")).lower() == normalized_email),
        {},
    )
    persona_name = str(resource.get("persona_name") or persona_id)
    viewer_name = str(resource.get("display_name") or email)
    last_updated = datetime.fromtimestamp(workbook_path.stat().st_mtime, timezone.utc).isoformat()
    kpis = [
        _build_kpi_detail(
            definition,
            user_measurements.get(str(definition["kpi_id"]), []),
            user_observations.get(str(definition["kpi_id"]), []),
            user_executions,
            period,
            last_updated,
            persona_id,
            persona_name,
            viewer_name,
            record_limit,
            record_offset,
        )
        for definition in definitions
        if definition.get("persona_id") in {persona_id, "all"}
    ]
    return sorted(kpis, key=_kpi_attention_order)


def _kpi_attention_order(kpi: dict[str, Any]) -> tuple[int, float, str]:
    status = str(kpi.get("status", "")).lower()
    if "below" in status or "above" in status:
        priority = 0
    elif status == "no data":
        priority = 3
    elif status == "no target configured":
        priority = 2
    else:
        priority = 1
    variance = kpi.get("variance_to_target")
    return priority, -abs(_number(variance)) if variance is not None else 0.0, str(kpi.get("kpi_name", ""))


def list_demo_responsibilities(
    persona_id: str,
    responsibility_id: str | None = None,
    status: str | None = None,
    workbook_path: Path = DEMO_WORKBOOK_PATH,
) -> dict[str, Any] | None:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Demo KPI workbook not found: {workbook_path}")
    responsibilities = [
        row for row in _read_demo_table("Responsibilities", workbook_path)
        if row.get("persona_id") == persona_id
    ]
    work_rows = [
        row for row in _read_demo_table("Responsibility_Work", workbook_path)
        if row.get("persona_id") == persona_id
    ]
    if responsibility_id:
        responsibilities = [row for row in responsibilities if row.get("responsibility_id") == responsibility_id]
        if not responsibilities:
            return None
    selected_status = status.casefold() if status else None
    result = []
    for item in responsibilities:
        linked_work = [
            row for row in work_rows
            if row.get("responsibility_id") == item.get("responsibility_id")
            and (selected_status is None or str(row.get("status", "")).casefold() == selected_status)
        ]
        counts: dict[str, int] = defaultdict(int)
        for row in work_rows:
            if row.get("responsibility_id") == item.get("responsibility_id"):
                counts[str(row.get("status", ""))] += 1
        result.append({
            **item,
            "status_counts": dict(counts),
            "work_items": sorted(linked_work, key=lambda row: str(row.get("updated_at", "")), reverse=True),
            "work_item_count": len(linked_work),
        })
    return {
        "responsibilities": result,
        "status_filter": status,
        "data_source": "DEMO / SYNTHETIC DATA",
    }


def list_demo_process_executions(
    email: str,
    status: str | None = None,
    search: str | None = None,
    process: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    limit: int = 25,
    offset: int = 0,
    workbook_path: Path = DEMO_WORKBOOK_PATH,
) -> dict[str, Any]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Demo KPI workbook not found: {workbook_path}")
    rows = _read_demo_table("Process_Executions", workbook_path)
    normalized_email = email.strip().lower()
    filtered = [
        row for row in rows
        if str(row.get("assigned_resource_email", "")).lower() == normalized_email
    ]
    if status:
        filtered = [
            row for row in filtered
            if str(row.get("process_execution_status", "")).casefold() == status.casefold()
        ]
    if search:
        query = search.strip().casefold()
        filtered = [
            row for row in filtered
            if query in " ".join(str(row.get(field, "")) for field in (
                "customer_request_id", "process_name", "task_name", "process_execution_status",
            )).casefold()
        ]
    if process:
        process_query = process.strip().casefold()
        filtered = [
            row for row in filtered
            if process_query in str(row.get("process_name", "")).casefold()
        ]
    if date_from:
        filtered = [
            row for row in filtered
            if row.get("process_start_timestamp")
            and date.fromisoformat(str(row["process_start_timestamp"])[:10]) >= date_from
        ]
    if date_to:
        filtered = [
            row for row in filtered
            if row.get("process_start_timestamp")
            and date.fromisoformat(str(row["process_start_timestamp"])[:10]) <= date_to
        ]
    now = datetime.now(timezone.utc)
    formatted = [_format_execution(row, now) for row in filtered]
    formatted.sort(key=lambda row: (row["started_at"], row["process_sequence"]), reverse=True)
    return {
        "items": formatted[offset : offset + limit],
        "total": len(formatted),
        "limit": limit,
        "offset": offset,
        "status_filter": status,
        "search": search,
        "process_filter": process,
        "date_from": date_from.isoformat() if date_from else None,
        "date_to": date_to.isoformat() if date_to else None,
    }


def _format_execution(row: dict[str, Any], now: datetime) -> dict[str, Any]:
    started = datetime.fromisoformat(row["process_start_timestamp"])
    ended = datetime.fromisoformat(row["process_end_timestamp"]) if row.get("process_end_timestamp") else None
    elapsed = round(((ended or now) - started).total_seconds() / 60, 1)
    return {
        "process_execution_id": row["process_execution_id"],
        "customer_request_id": row["customer_request_id"],
        "process_id": row["process_id"],
        "process_name": row["process_name"],
        "activity_name": row["activity_name"],
        "task_name": row["task_name"],
        "process_sequence": int(row["process_sequence"]),
        "status": row["process_execution_status"],
        "outcome": row["execution_outcome"],
        "persona_id": row["persona_id"],
        "resource_id": row["assigned_resource_id"],
        "started_at": row["process_start_timestamp"],
        "ended_at": row.get("process_end_timestamp"),
        "duration_minutes": _number(row.get("actual_duration_minutes")) if row.get("actual_duration_minutes") is not None else None,
        "elapsed_minutes": elapsed if ended is None and row["process_execution_status"] == "In Progress" else None,
        "sla_target_minutes": _number(row.get("sla_target_duration_minutes")),
        "retry_count": int(row.get("retry_count") or 0),
        "failure_reason": row.get("failure_reason"),
        "alert_id": row.get("alert_id"),
    }


def build_individual_dashboard_payload(
    email: str,
    viewer: dict[str, str],
    persona_id: str,
    workbook_path: Path = DEMO_WORKBOOK_PATH,
    include_process_executions: bool = True,
) -> dict[str, Any]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Demo KPI workbook not found: {workbook_path}")

    definitions = _read_demo_table("KPI_Definitions", workbook_path)
    measurements = _read_demo_table("KPI_Measurements", workbook_path)
    all_executions = _read_demo_table("Process_Executions", workbook_path)
    alerts = _read_demo_table("Alert_History", workbook_path)
    requests = _read_demo_table("Customer_Requests", workbook_path)
    responsibilities = _read_demo_table("Responsibilities", workbook_path)
    responsibility_work = _read_demo_table("Responsibility_Work", workbook_path)

    normalized_email = email.strip().lower()
    executions = [
        row for row in all_executions
        if str(row.get("assigned_resource_email", "")).lower() == normalized_email
    ]
    definitions_by_id = {row["kpi_id"]: row for row in definitions}
    measurements_by_id: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in measurements:
        if str(row.get("resource_email", "")).lower() == normalized_email:
            measurements_by_id[row["kpi_id"]].append(row)

    kpis: list[dict[str, Any]] = []
    for definition in definitions:
        role_id = definition["persona_id"]
        if role_id not in {persona_id, "all"}:
            continue
        metric_id = definition["kpi_id"]
        if definition["aggregation_method"] == "derived_from_executions":
            relevant = [row for row in executions if row["process_execution_status"] != "Pending"]
            completed = [row for row in relevant if row["process_execution_status"] == "Completed"]
            if metric_id == "process-completion-rate":
                denominator = len(relevant)
                value = round(len(completed) / denominator * 100, 2) if denominator else 0.0
            elif metric_id == "average-process-duration":
                durations = [_number(row["actual_duration_minutes"]) for row in completed if row.get("actual_duration_minutes") is not None]
                value = round(sum(durations) / len(durations), 2) if durations else 0.0
            else:
                denominator = len(completed)
                on_time = [
                    row for row in completed
                    if _number(row.get("actual_duration_minutes")) <= _number(row.get("sla_target_duration_minutes"))
                ]
                value = round(len(on_time) / denominator * 100, 2) if denominator else 0.0
            kpis.append({
                "id": metric_id,
                "name": definition["kpi_name"],
                "description": definition["description"],
                "value": value,
                "unit": definition["measurement_unit"],
                "period": "All demo executions",
                "status": "DEMO / illustrative",
                "source": definition["source"],
                "sample_size": len(completed) if metric_id == "average-process-duration" else len(relevant),
                "series": [],
            })
            continue

        points = sorted(measurements_by_id.get(metric_id, []), key=lambda row: row["period_end"])
        if not points:
            continue
        series = [
            {
                "time": str(point["period_end"])[:10],
                "value": _measurement_value(point, definition),
            }
            for point in points
        ]
        latest = points[-1]
        kpis.append({
            "id": metric_id,
            "name": definition["kpi_name"],
            "description": definition["description"],
            "value": series[-1]["value"],
            "target": _number(definition.get("target_value")) if definition.get("target_value") is not None else None,
            "target_source": definition.get("target_source"),
            "unit": definition["measurement_unit"],
            "period": "Weekly",
            "status": (
                f"{_status_for_target(series[-1]['value'], _number(definition.get('target_value')), definition.get('direction'))} (illustrative)"
                if definition.get("target_source") and definition.get("target_value") is not None
                else "DEMO / illustrative"
            ),
            "source": definition["source"],
            "sample_size": int(latest.get("sample_size") or 0),
            "series": series,
        })

    now = datetime.now(timezone.utc)
    formatted_executions = [_format_execution(row, now) for row in executions]
    formatted_executions.sort(key=lambda row: (row["started_at"], row["process_sequence"]))
    scoped_request_ids = {row["customer_request_id"] for row in executions}
    scoped_requests = [row for row in requests if row["request_id"] in scoped_request_ids]
    scoped_alerts = [row for row in alerts if row["customer_request_id"] in scoped_request_ids]
    scoped_work = [row for row in responsibility_work if row.get("persona_id") == persona_id]
    scoped_responsibilities = [row for row in responsibilities if row.get("persona_id") == persona_id]
    completed = [row for row in executions if row["process_execution_status"] == "Completed"]
    active = [row for row in executions if row["process_execution_status"] == "In Progress"]
    overdue_request_count = sum(1 for row in scoped_requests if row["request_status"] == "Overdue")
    active_request_ids = {
        row["customer_request_id"] for row in formatted_executions
        if row["status"] == "In Progress"
    }
    workflow_request = max(
        (row for row in formatted_executions if row["customer_request_id"] in active_request_ids),
        key=lambda row: row["started_at"],
        default=max(formatted_executions, key=lambda row: row["started_at"], default=None),
    )
    current_workflow = (
        sorted(
            [
                row for row in formatted_executions
                if row["customer_request_id"] == workflow_request["customer_request_id"]
            ],
            key=lambda row: row["process_sequence"],
        )
        if workflow_request else []
    )

    persona_name = next(
        (row["persona_name"] for row in _read_persona_rows() if row["persona_id"] == persona_id),
        persona_id,
    )
    role_number = int(persona_id.removeprefix("role-"))
    kpis.sort(
        key=lambda item: (
            0 if "below" in str(item.get("status", "")).lower() or "above" in str(item.get("status", "")).lower()
            else 1 if "illustrative)" in str(item.get("status", "")).lower()
            else 2,
        )
    )
    work_statuses = ("Pending", "In Progress", "Completed", "Blocked", "Overdue")
    work_status_counts = {
        status: sum(1 for row in scoped_work if row.get("status") == status)
        for status in work_statuses
    }
    status_items = (
        [
            {"label": status, "value": count, "status_filter": status}
            for status, count in work_status_counts.items()
        ]
        if scoped_work
        else [
            {"label": "Customer requests", "value": len(scoped_requests)},
            {"label": "Processes in progress", "value": len(active)},
            {"label": "Completed processes", "value": len(completed)},
            {"label": "Overdue requests", "value": overdue_request_count},
            {"label": "Open alerts", "value": sum(1 for row in scoped_alerts if row["resolution_status"] == "Open")},
        ]
    )
    responsibility_items = []
    for responsibility in scoped_responsibilities:
        linked = [row for row in scoped_work if row.get("responsibility_id") == responsibility.get("responsibility_id")]
        counts = {
            status: sum(1 for row in linked if row.get("status") == status)
            for status in work_statuses
        }
        responsibility_items.append({
            **responsibility,
            "status_counts": counts,
            "work_item_count": len(linked),
        })
    action_items = [
        {
            "title": f"Review {item['name']}",
            "priority": item["status"],
            "kpi": item["name"],
        }
        for item in kpis if "below" in str(item.get("status", "")).lower()
    ]
    return {
        "persona": persona_id,
        "viewer": {**viewer, "designation": f"Role {role_number:02d} – {persona_name}"},
        "data_source": "DEMO / SYNTHETIC DATA",
        "updated_at": int(now.timestamp() * 1000),
        "status_panel": {
            "title": "My process performance",
            "summary": "Illustrative records scoped to this authenticated demo account.",
            "items": status_items,
        },
        "metric_panel": {
            "title": "Individual KPI scorecard",
            "items": [
                {"label": item["name"], "value": item["value"], "unit": item["unit"], "status": item["status"]}
                for item in kpis
            ],
        },
        "responsibility_panel": {"title": "Responsibilities", "items": responsibility_items},
        "action_panel": {"title": "KPI priorities", "items": action_items},
        "workflow_panel": {
            "title": "Current request workflow",
            "request_id": workflow_request["customer_request_id"] if workflow_request else None,
            "items": [
                {"step": row["process_name"], "status": row["status"]}
                for row in current_workflow
            ],
        },
        "kpis": kpis,
        "requests": scoped_requests,
        "process_executions": formatted_executions if include_process_executions else [],
        "alerts": scoped_alerts,
    }


def _read_persona_rows() -> list[dict[str, Any]]:
    if not DEMO_WORKBOOK_PATH.exists():
        return []
    return _read_demo_table("Persona_Resource_Mapping")
