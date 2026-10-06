from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
    ".md": "markdown",
    ".csv": "csv",
    ".xlsx": "xlsx",
    ".pptx": "pptx",
    ".html": "html",
    ".json": "json",
}


def compute_content_hash(file_path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(file_path, "rb") as handle:
        for chunk in iter(lambda: handle.read(8192), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_text_file(path: str | Path) -> str:
    return Path(path).read_text(encoding="utf-8", errors="replace")


def extract_pdf_text(path: str | Path) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"\n--- Page {page_number} ---\n{text}")
    return "\n".join(pages)


def extract_docx_text(path: str | Path) -> str:
    doc = DocxDocument(str(path))
    paragraphs = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs)


def extract_xlsx_text(path: str | Path) -> str:
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    rows: list[str] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows.append(f"\n--- Sheet: {sheet_name} ---\n")
        for row in sheet.iter_rows(values_only=True):
            non_empty = [str(cell) for cell in row if cell is not None]
            if non_empty:
                rows.append(" | ".join(non_empty))
    return "\n".join(rows)


def extract_csv_text(path: str | Path) -> str:
    rows: list[str] = []
    with open(path, newline="", encoding="utf-8", errors="replace") as handle:
        reader = csv.reader(handle)
        for row in reader:
            if any(cell.strip() for cell in row):
                rows.append(" | ".join(row))
    return "\n".join(rows)


def extract_pptx_text(path: str | Path) -> str:
    prs = Presentation(str(path))
    slides: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        pieces = [shape.text for shape in slide.shapes if getattr(shape, "text", "").strip()]
        slides.append(f"\n--- Slide {index} ---\n" + "\n".join(pieces))
    return "\n".join(slides)


def extract_html_text(path: str | Path) -> str:
    html = read_text_file(path)
    soup = BeautifulSoup(html, "html.parser")
    return soup.get_text(separator="\n", strip=True)


def extract_json_text(path: str | Path) -> str:
    data = json.loads(read_text_file(path))
    return json.dumps(data, ensure_ascii=False, indent=2)


def extract_uploaded_text(path: str | Path, extension: str) -> tuple[str, dict[str, Any]]:
    ext = extension.lower()
    if ext == ".pdf":
        text = extract_pdf_text(path)
    elif ext == ".docx":
        text = extract_docx_text(path)
    elif ext in {".txt", ".md"}:
        text = read_text_file(path)
    elif ext == ".csv":
        text = extract_csv_text(path)
    elif ext == ".xlsx":
        text = extract_xlsx_text(path)
    elif ext == ".pptx":
        text = extract_pptx_text(path)
    elif ext == ".html":
        text = extract_html_text(path)
    elif ext == ".json":
        text = extract_json_text(path)
    else:
        raise ValueError(f"Unsupported extension: {extension}")

    metadata = {"source_extension": ext, "char_count": len(text), "word_count": len(text.split())}
    return text, metadata


def is_supported_extension(extension: str) -> bool:
    return extension.lower() in SUPPORTED_EXTENSIONS
